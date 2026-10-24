import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy.orm import sessionmaker
from sqlalchemy import select

from korsik.auth import Actor, get_actor
from korsik.db import get_db, make_engine
from korsik.models import Base, Entity, Job, User
from korsik import knowledge, retrieval
from korsik.common import save_entity
from korsik.parsing import block, digest, parse_document


def test_parser_storage_failure_releases_admission_for_two_attempts(monkeypatch):
    import asyncio
    import errno
    import importlib.util
    import sys
    from pathlib import Path
    from fastapi import HTTPException
    from korsik import parsing

    monkeypatch.setitem(sys.modules, "parsing", parsing)
    path = Path(__file__).resolve().parents[1] / "services/model_server/parser_app.py"
    spec = importlib.util.spec_from_file_location("storage_failure_parser", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    def full_storage(**kwargs):
        raise OSError(errno.ENOSPC, "synthetic temporary storage full")

    monkeypatch.setattr(module.tempfile, "TemporaryDirectory", full_storage)
    for _ in range(2):
        with pytest.raises(HTTPException) as raised:
            asyncio.run(module.parse(file=None, skip_pages="[]", page_options="{}"))
        assert raised.value.status_code == 507
        assert raised.value.detail == "parser_storage_full"
    assert module.BUSY.acquire(blocking=False)
    module.BUSY.release()


@pytest.fixture
def kb(tmp_path, monkeypatch):
    engine = make_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = sessionmaker(engine, expire_on_commit=False)()
    actor = Actor("admin", "HR", "HR_ADMIN", {"all": True})
    db.add(User(id=actor.id, email="admin@kb.test", password_hash="not-used", name=actor.name, role=actor.role, scope=actor.scope, enabled=True))
    db.commit()
    app = FastAPI()
    app.include_router(knowledge.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_actor] = lambda: actor
    monkeypatch.setattr(knowledge, "DATA", tmp_path)
    monkeypatch.setattr(retrieval, "MODE", "CORE")
    monkeypatch.setattr(retrieval, "_request", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("offline")))
    yield TestClient(app), db, actor
    db.close()
    engine.dispose()


def create_text(client, db, text="Обед доступен с 12:00 до 13:00.", site="a"):
    doc = client.post("/knowledge/documents", json={"title": "Правила", "scope": {"site_ids": [site], "employment_stages": ["ACTIVE"]}, "document_category": "REGULATION"}).json()
    response = client.post(f"/knowledge/documents/{doc['id']}/revisions", files={"file": ("rules.txt", text.encode(), "text/plain")}, data={"reason": "Обновление правил"}, headers={"If-Match": doc["etag"]})
    assert response.status_code == 202, response.text
    revision = response.json()["revision"]
    job = db.scalar(select(Job).where(Job.kind == "KNOWLEDGE_INGEST", Job.payload.is_not(None)).order_by(Job.created_at.desc()))
    knowledge.process_ingest_job(db, job)
    return doc, client.get(f"/knowledge/revisions/{revision['id']}").json()


def page_review(client, revision_id):
    pages = client.get(f"/knowledge/revisions/{revision_id}/pages")
    response = client.patch(f"/knowledge/revisions/{revision_id}/pages/1/review", json={"action": "REVIEW_CONTENT", "reason": "Сверено с исходным документом", "expected_extraction_revision_id": pages.json()["extraction_revision_id"]}, headers={"If-Match": pages.headers["etag"]})
    assert response.status_code == 200, response.text


def action(client, revision_id, path, payload=None):
    current = client.get(f"/knowledge/revisions/{revision_id}")
    result = client.post(f"/knowledge/revisions/{revision_id}/{path}", json=payload or {}, headers={"If-Match": current.headers["etag"]})
    if path in {"qa-runs", "candidate-build"} and result.status_code == 202:
        import httpx
        from korsik.common import serialize
        db = client.app.dependency_overrides[get_db]()
        op_id = result.json()["id"]
        job = db.scalar(select(Job).where(Job.kind==("KNOWLEDGE_QA" if path=="qa-runs" else "KNOWLEDGE_CANDIDATE_BUILD")).order_by(Job.created_at.desc()))
        knowledge.process_ingest_job(db, job)
        return httpx.Response(202, json=serialize(db.get(Entity, op_id)))
    return result


def publish(client, revision_id):
    page_review(client, revision_id)
    assert action(client, revision_id, "candidate-build").status_code == 202
    assert action(client, revision_id, "qa-runs").status_code == 202
    approve = action(client, revision_id, "review", {"action": "APPROVE", "reason": "Проверено HR"})
    assert approve.status_code == 200, approve.text
    result = action(client, revision_id, "publish", {"reason": "Утверждено"})
    assert result.status_code == 202, result.text


def test_upload_is_not_searchable_and_human_coverage_required(kb):
    client, db, actor = kb
    doc, revision = create_text(client, db)
    assert revision["status"] == "REVIEW_REQUIRED"
    employee = save_entity(db, "employee", {"site_id": "a", "employment_stage": "ACTIVE", "employment_type": "OFFICE"})
    assert retrieval.published_corpus(db, employee) == {}
    assert action(client, revision["id"], "candidate-build").status_code == 202
    qa = action(client, revision["id"], "qa-runs").json()
    assert qa["qa_results"][0]["status"] == "FAILED"
    assert action(client, revision["id"], "review", {"action": "APPROVE", "reason": "Проверено"}).status_code == 409
    publish(client, revision["id"])
    assert revision["id"] in retrieval.published_corpus(db, employee)


def test_published_extraction_immutable_and_stale_etag_denied(kb):
    client, db, _ = kb
    _, revision = create_text(client, db)
    extraction = client.get(f"/knowledge/revisions/{revision['id']}/extraction")
    b = extraction.json()["blocks"][0]
    payload = {"patch_type": "TEXT", "text": "Исправленный текст", "reason": "Проверка", "expected_extraction_revision_id": extraction.json()["extraction_revision_id"]}
    path = f"/knowledge/revisions/{revision['id']}/extraction/blocks/{b['id']}"
    assert client.patch(path, json=payload, headers={"If-Match": extraction.headers["etag"]}).status_code == 200
    assert client.patch(path, json=payload, headers={"If-Match": extraction.headers["etag"]}).status_code == 412
    publish(client, revision["id"])
    latest = client.get(f"/knowledge/revisions/{revision['id']}/extraction")
    assert client.patch(path, json=payload, headers={"If-Match": latest.headers["etag"]}).status_code == 409


def test_replayed_scheduled_publish_keeps_published_revision(kb):
    client, db, actor = kb
    doc, revision = create_text(client, db)
    publish(client, revision["id"])
    replay = Job(kind="KNOWLEDGE_PUBLISH", unique_key="publish-replay-test", payload={"revision_id": revision["id"], "expected_snapshot_id": "obsolete"})
    db.add(replay)
    db.flush()
    assert knowledge.process_ingest_job(db, replay)
    assert db.get(Entity, revision["id"]).status == "PUBLISHED"
    assert db.get(Entity, doc["id"]).data["active_revision_id"] == revision["id"]


@pytest.mark.parametrize("change", [None, "disabled", "role", "scope", "review", "postwrite_fault"])
def test_scheduled_publish_has_terminal_operation_and_fresh_authority(kb, change, monkeypatch):
    from datetime import timedelta
    from korsik.db import now

    client, db, actor = kb
    doc, revision = create_text(client, db)
    page_review(client, revision["id"])
    assert action(client, revision["id"], "candidate-build").status_code == 202
    assert action(client, revision["id"], "qa-runs").status_code == 202
    assert action(client, revision["id"], "review", {"action": "APPROVE", "reason": "Проверено HR"}).status_code == 200
    body = {"effective_at": (now() + timedelta(hours=1)).isoformat(), "reason": "Плановая публикация"}
    scheduled = action(client, revision["id"], "publish", body)
    assert scheduled.status_code == 202
    operation_id = scheduled.json()["id"]
    duplicate = action(client, revision["id"], "publish", body)
    assert duplicate.status_code == 202 and duplicate.json()["id"] == operation_id
    jobs = db.scalars(select(Job).where(Job.kind == "KNOWLEDGE_PUBLISH")).all()
    assert len(jobs) == 1 and jobs[0].payload["operation_id"] == operation_id
    job = jobs[0]
    job.status = "PROCESSING"
    job.run_at = now() - timedelta(seconds=1)
    job.lease_until = now() + timedelta(minutes=10)
    job.payload = {**job.payload, "_lease_token": "claimed-publish"}
    user = db.get(User, actor.id)
    current = db.get(Entity, revision["id"])
    if change == "disabled":
        user.enabled = False
    elif change == "role":
        user.role = "HR_SPECIALIST"
    elif change == "scope":
        user.scope = {"site_ids": ["other"]}
    elif change == "review":
        current.data = {**current.data, "review_fingerprint": "changed-review"}
    elif change == "postwrite_fault":
        from fastapi import HTTPException
        original = knowledge._publish
        def failed_after_write(*args, **kwargs):
            original(*args, **kwargs)
            raise HTTPException(409, "injected_after_write")
        monkeypatch.setattr(knowledge, "_publish", failed_after_write)
    old_snapshot_id = db.scalar(select(Entity).where(Entity.kind == "knowledge_head")).data["snapshot_id"]
    db.commit()
    assert knowledge.process_ingest_job(db, job)
    db.commit()
    op = db.get(Entity, operation_id)
    assert op.status == ("SUCCEEDED" if change is None else "FAILED")
    assert op.data.get("published") is (change is None)
    if change is None:
        assert db.get(Entity, revision["id"]).status == "PUBLISHED"
        snapshot = op.data["kb_snapshot_id"]
        assert knowledge.process_ingest_job(db, job)
        assert db.get(Entity, operation_id).data["kb_snapshot_id"] == snapshot
    else:
        assert not db.get(Entity, doc["id"]).data.get("active_revision_id")
        assert db.get(Entity, revision["id"]).status == "REVIEW_REQUIRED"
        assert db.scalar(select(Entity).where(Entity.kind == "knowledge_head")).data["snapshot_id"] == old_snapshot_id


def test_revision_comparison_preserves_pins_and_compares_physical_structure(kb):
    client, db, actor = kb
    doc, first = create_text(client, db, "Исходный текст")
    old = db.get(Entity, first["id"])
    old_ext = knowledge._new_extraction(db, old, {"pages": [{"page_no": n, "status": "PROCESSED"} for n in (1, 2)],
        "blocks": [block("TEXT", "Не допускается передавать без согласования. Срок 2 дня.", 1),
                   block("TABLE", "A | 1", 1, grid=[["A", "1"]], reading_order=1),
                   block("FIGURE", "", 1, image_sha256="a"*64, reading_order=2),
                   block("TEXT", "Прежняя вторая страница", 2),
                   block("TABLE", "B | old", 2, grid=[["B", "old"]], reading_order=1),
                   block("FIGURE", "", 2, image_sha256="c"*64, reading_order=2)]})
    db.commit()
    current_doc = client.get(f"/knowledge/documents/{doc['id']}")
    assert client.patch(f"/knowledge/documents/{doc['id']}", json={"title": "Новое название"}, headers={"If-Match": current_doc.headers["etag"]}).status_code == 200
    current_doc = client.get(f"/knowledge/documents/{doc['id']}")
    posted = client.post(f"/knowledge/documents/{doc['id']}/revisions/article", json={"text": "Новая редакция", "reason": "Обновление"}, headers={"If-Match": current_doc.headers["etag"]})
    assert posted.status_code == 201
    target = db.get(Entity, posted.json()["id"])
    new_ext = knowledge._new_extraction(db, target, {"pages": [{"page_no": n, "status": "PROCESSED"} for n in (1, 3)],
        "blocks": [block("TEXT", "Допускается передавать после согласования. Срок 3 дня.", 1),
                   block("TABLE", "A | 2", 1, grid=[["A", "2"]], reading_order=1),
                   block("FIGURE", "", 1, image_sha256="b"*64, reading_order=2),
                   block("TEXT", "Новая третья страница", 3),
                   block("TABLE", "C | new", 3, grid=[["C", "new"]], reading_order=1),
                   block("FIGURE", "", 3, image_sha256="d"*64, reading_order=2)]})
    db.commit()
    route = f"/knowledge/revisions/{target.id}/compare"
    params = {"base_revision_id": old.id, "extraction_revision_id": new_ext.id, "base_extraction_revision_id": old_ext.id}
    response = client.get(route, params=params)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["base"]["extraction_revision_id"] == old_ext.id
    assert result["target"]["extraction_revision_id"] == new_ext.id
    assert result["comparison"]["semantic_diff"] is False
    summary = result["summary"]
    assert summary["pages_added"] == 1 and summary["pages_removed"] == 1 and summary["pages_changed"] == 1
    assert summary["tables_changed"] == 1 and summary["figures_changed"] == 1
    assert summary["tables_added"] == summary["tables_removed"] == summary["figures_added"] == summary["figures_removed"] == 1
    assert result["pages"][0]["text"]["changes"][0]["contains_numbers"]
    assert result["pages"][0]["text"]["changes"][0]["contains_negation"]
    title = next(row for row in result["metadata"] if row["field"] == "title")
    assert title["before"] == "Правила" and title["after"] == "Новое название"
    knowledge._new_extraction(db, target, {"pages": [{"page_no": 1, "status": "PROCESSED"}], "blocks": [block("TEXT", "Более позднее исправление", 1)]})
    db.commit()
    assert client.get(route, params=params).json()["pages"] == result["pages"]
    # A pin from another revision is not silently redirected to its latest state.
    assert client.get(route, params={**params, "base_extraction_revision_id": new_ext.id}).status_code == 404
    # A failed extraction is unavailable, not proof that earlier content vanished.
    failed = knowledge._new_extraction(db, target, {"pages": [{"page_no": 1, "status": "FAILED"}], "blocks": []})
    db.commit()
    incomplete = client.get(route, params={**params, "extraction_revision_id": failed.id}).json()
    assert incomplete["pages"][0]["status"] == "UNAVAILABLE" and incomplete["pages"][0]["text"]["status"] == "UNAVAILABLE"


def test_revision_comparison_limits_scope_and_unavailable_history(kb):
    import json
    client, db, actor = kb
    doc, first = create_text(client, db)
    revision = db.get(Entity, first["id"])
    revision.data = {key: value for key, value in revision.data.items() if key != "metadata_snapshot"}
    long_text = "Правило 123 не меняется.\n"*1000
    knowledge._new_extraction(db, revision, {"pages": [{"page_no": n, "status": "PROCESSED"} for n in range(1, 80)],
        "blocks": [block("TEXT", long_text, n) for n in range(1, 80)]})
    db.commit()
    route = f"/knowledge/revisions/{revision.id}/compare"
    result = client.get(route, params={"base_revision_id": revision.id}).json()
    assert result["truncated"] and "PHYSICAL_PAGE_LIMIT" in result["truncation_reasons"] and "TEXT_CHAR_LIMIT" in result["truncation_reasons"]
    assert len(json.dumps(result, ensure_ascii=False).encode()) <= result["limits"]["response_bytes"]
    assert result["metadata_history_unavailable"] == ["title", "scope", "document_category"]
    assert result["current_document_metadata"]["is_historical"] is False
    _, foreign = create_text(client, db, site="b")
    assert client.get(route, params={"base_revision_id": foreign["id"]}).status_code == 422
    client.app.dependency_overrides[get_actor] = lambda: Actor(actor.id, "Scoped", "HR_ADMIN", {"site_ids": ["a"]})
    assert client.get(route, params={"base_revision_id": foreign["id"]}).status_code == 404
    doc_entity = db.get(Entity, doc["id"])
    doc_entity.data = {**doc_entity.data, "source_revoked": True}
    db.commit()
    assert client.get(route, params={"base_revision_id": revision.id}).status_code == 404


def test_article_dates_reject_invalid_or_inverted_window(kb):
    client, db, actor = kb
    doc = client.post("/knowledge/documents", json={"title": "Статья", "scope": {"site_ids": ["a"]}}).json()
    route = f"/knowledge/documents/{doc['id']}/revisions/article"
    for dates in [{"effective_from": "bad-date"}, {"effective_from": "2030-01-01", "effective_to": "2020-01-01"}]:
        result = client.post(route, json={"text": "Правила компании", **dates}, headers={"If-Match": doc["etag"]})
        assert result.status_code == 422
    assert not db.scalar(select(Entity).where(Entity.kind=="document_revision", Entity.parent_id==doc["id"]))


def test_identical_upload_can_be_explicitly_versioned(kb):
    client, db, actor = kb
    doc, revision = create_text(client, db)
    path = f"/knowledge/documents/{doc['id']}/revisions"
    current = client.get(f"/knowledge/documents/{doc['id']}")
    def upload(data):
        return client.post(path, files={"file":("rules.txt","Обед доступен с 12:00 до 13:00.".encode(),"text/plain")}, data=data, headers={"If-Match":current.headers['etag']})
    repeated=upload({"reason":"Повторная загрузка"})
    assert repeated.status_code==202 and repeated.json()['duplicate'] is True
    assert repeated.json()['revision']['id']==revision['id']
    new=upload({"reason":"Повторное утверждение того же текста", "force_new_version":"true"})
    assert new.status_code==202 and new.json()['revision']['id']!=revision['id']
    assert new.json()['revision']['status']=='UPLOADED'


def test_plural_legal_scope_and_site_legal_grants_are_conjunction(kb):
    client, db, actor = kb
    actor.scope={"site_ids":["a"],"legal_entity_ids":["org"]}
    allowed={"title":"Инструкция из пробела", "scope":{"site_ids":["a"],"legal_entity_ids":["org"],"employment_stages":["ACTIVE"]}}
    response=client.post('/knowledge/documents',json=allowed)
    assert response.status_code==201,response.text
    assert response.json()['scope']['legal_entity_ids']==['org']
    for scope in [{"site_ids":["b"],"legal_entity_ids":["org"]},{"site_ids":["a"],"legal_entity_ids":["other"]}]:
        assert client.post('/knowledge/documents',json={"title":"Вне области","scope":scope}).status_code==403
    assert knowledge._scope_subset({'legal_entity_ids':['org']},{'legal_entity_id':'org'})
    assert not knowledge._scope_subset({'legal_entity_ids':['org','other']},{'legal_entity_id':'org'})


def test_qa_is_durable_and_cancelled_or_stale_jobs_create_no_artifacts(kb, monkeypatch):
    client, db, actor = kb
    doc, revision = create_text(client, db)
    page_review(client, revision['id'])
    action(client,revision['id'],'candidate-build')
    from korsik import knowledge_admin
    called=[]
    original=knowledge_admin.evaluate_control_suite
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',lambda *a,**k:called.append(True))
    def submit():
        current=client.get(f"/knowledge/revisions/{revision['id']}")
        result=client.post(f"/knowledge/revisions/{revision['id']}/qa-runs",json={},headers={'If-Match':current.headers['etag']})
        assert result.status_code==202 and result.json()['status']=='PENDING'
        return db.get(Entity,result.json()['id']),db.scalar(select(Job).where(Job.kind=='KNOWLEDGE_QA').order_by(Job.created_at.desc()))
    operation,job=submit()
    assert not called and not db.get(Entity,revision['id']).data.get('qa_artifact_id')
    operation.status='CANCELLED';operation.data={**operation.data,'cancel_requested':True};db.commit()
    assert knowledge.process_ingest_job(db,job)
    assert operation.status=='CANCELLED' and not called
    operation,job=submit()
    save_entity(db,'knowledge_control_question',{'current_version_id':'new-pinned-version'},status='ACTIVE');db.commit()
    assert knowledge.process_ingest_job(db,job)
    assert operation.status=='FAILED' and operation.data['error']=='qa_controls_or_model_profile_changed'
    assert not called and not db.get(Entity,revision['id']).data.get('qa_artifact_id')
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',original)


def test_qa_lease_loss_discards_old_worker_and_successor_can_finish(kb,monkeypatch):
    client,db,actor=kb
    doc,revision=create_text(client,db);page_review(client,revision['id']);action(client,revision['id'],'candidate-build')
    current=client.get(f"/knowledge/revisions/{revision['id']}")
    queued=client.post(f"/knowledge/revisions/{revision['id']}/qa-runs",json={},headers={'If-Match':current.headers['etag']})
    operation=db.get(Entity,queued.json()['id'])
    job=db.scalar(select(Job).where(Job.kind=='KNOWLEDGE_QA').order_by(Job.created_at.desc()))
    from datetime import timedelta
    job.payload={**job.payload,'_lease_token':'old-worker'}
    job.status='PROCESSING';job.lease_until=knowledge.now()+timedelta(seconds=300);db.commit()
    from korsik import knowledge_admin
    original=knowledge_admin.evaluate_control_suite
    def superseded(db,candidate,actor,checkpoint=None):
        job.payload={**job.payload,'_lease_token':'successor-worker'};db.commit()
        checkpoint(0)
        raise AssertionError('old worker must stop before result')
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',superseded)
    knowledge.process_ingest_job(db,job)
    assert operation.status=='RUNNING' and not db.get(Entity,revision['id']).data.get('qa_artifact_id')
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',original)
    knowledge.process_ingest_job(db,job)
    assert operation.status=='SUCCEEDED'
    assert db.get(Entity,revision['id']).data['qa_artifact_id']


def test_draft_scope_and_source_are_not_employee_accessible(kb):
    client, db, actor = kb
    doc, revision = create_text(client, db, site="a")
    employee = save_entity(db, "employee", {"site_id": "b", "employment_stage": "ACTIVE"})
    actor.role, actor.employee_id = "EMPLOYEE", employee.id
    assert client.get(f"/knowledge/documents/{doc['id']}").status_code == 404
    assert client.get(f"/knowledge/revisions/{revision['id']}/source").status_code == 404
    assert client.get("/knowledge/documents").json()["items"] == []
    assert client.get("/knowledge/snapshots/current").status_code == 403


def test_atomic_bundle_rejects_missing_member_then_publishes_both(kb):
    client, db, _ = kb
    doc, file_revision = create_text(client, db)
    page_review(client, file_revision["id"])
    source = client.get(f"/knowledge/revisions/{file_revision['id']}/extraction").json()
    source_block = source["blocks"][0]
    article_doc = client.post("/knowledge/documents", json={"title": "Обед", "scope": {"site_ids": ["a"], "employment_stages": ["ACTIVE"]}, "document_category": "INSTRUCTION"}).json()
    article = client.post(f"/knowledge/documents/{article_doc['id']}/revisions/article", headers={"If-Match": article_doc["etag"]}, json={"text": "Обед с 12:00 до 13:00.", "dependencies": [{"source_revision_id": file_revision["id"], "source_extraction_revision_id": source["extraction_revision_id"], "source_block_ids": [source_block["id"]], "source_block_sha256": {source_block["id"]: digest(source_block["text"].encode())}, "rationale": "Основание времени обеда"}]}).json()
    page_review(client, article["id"])
    current_article = client.get(f"/knowledge/revisions/{article['id']}").json()
    built = action(client, file_revision["id"], "candidate-build", {"additional_revisions": [{"revision_id": article["id"], "expected_extraction_revision_id": current_article["latest_extraction_revision_id"], "expected_content_fingerprint": current_article["content_fingerprint"]}]})
    assert built.status_code == 202, built.text
    assert action(client, file_revision["id"], "qa-runs").status_code == 202
    for id in [article["id"], file_revision["id"]]:
        assert action(client, id, "review", {"action": "APPROVE", "reason": "Сверено с источником"}).status_code == 200
    missing = action(client, file_revision["id"], "publish", {})
    assert missing.status_code == 422
    assert client.get(f"/knowledge/documents/{doc['id']}").json()["active_revision_id"] is None
    current_article = client.get(f"/knowledge/revisions/{article['id']}").json()
    result = action(client, file_revision["id"], "publish", {"additional_revision_reviews": [{"revision_id": article["id"], "expected_review_fingerprint": current_article["review_fingerprint"]}]})
    assert result.status_code == 202, result.text
    assert client.get(f"/knowledge/documents/{doc['id']}").json()["active_revision_id"] == file_revision["id"]
    assert client.get(f"/knowledge/documents/{article_doc['id']}").json()["active_revision_id"] == article["id"]


def test_partial_failure_checkpoint_and_retry_preserves_reviewed_pages(kb, monkeypatch):
    client, db, _ = kb
    def broken(path, output, checkpoint):
        checkpoint([{"page_no": 1, "status": "PROCESSED", "review_status": "PENDING", "issues": []}, {"page_no": 2, "status": "FAILED", "review_status": "PENDING", "issues": ["PARSE_ERROR"]}], [block("TEXT", "Первоначальный текст", 1)])
        raise RuntimeError("page decoder failed")
    monkeypatch.setattr(knowledge, "_run_parser", broken)
    _, revision = create_text(client, db)
    assert revision["status"] == "FAILED"
    extraction = client.get(f"/knowledge/revisions/{revision['id']}/extraction")
    block_id = extraction.json()["blocks"][0]["id"]
    response = client.patch(f"/knowledge/revisions/{revision['id']}/extraction/blocks/{block_id}", json={"patch_type": "TEXT", "text": "Исправлено человеком", "reason": "OCR ошибка"}, headers={"If-Match": extraction.headers["etag"]})
    assert response.status_code == 200
    def success(path, output, checkpoint, *, skip_pages=None):
        assert skip_pages == [1]
        checkpoint([{"page_no": n, "status": "PROCESSED", "review_status": "PENDING", "issues": []} for n in [1, 2]], [block("TEXT", "Перезаписанный машиной текст", 1), block("TEXT", "Вторая страница", 2)])
        return {"parser_version": "test"}
    monkeypatch.setattr(knowledge, "_run_parser", success)
    assert action(client, revision["id"], "retry", {}).status_code == 202
    job = db.scalar(select(Job).where(Job.kind == "KNOWLEDGE_INGEST").order_by(Job.created_at.desc()))
    knowledge.process_ingest_job(db, job)
    latest = client.get(f"/knowledge/revisions/{revision['id']}/extraction").json()
    assert latest["blocks"][0]["text"] == "Исправлено человеком"
    assert latest["blocks"][0]["id"] == block_id
    assert latest["pages"][1]["status"] == "PROCESSED"


def test_pptx_tables_remain_structured_and_require_review(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(2)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Условие", "Время"
    table.cell(1, 0).text, table.cell(1, 1).text = "Обед", "12:00"
    path = tmp_path / "sample.pptx"
    presentation.save(path)
    result = parse_document(path, tmp_path / "previews")
    table_block = next(b for b in result["blocks"] if b["kind"] == "TABLE")
    assert table_block["grid"] == [["Условие", "Время"], ["Обед", "12:00"]]
    assert "TABLE_STRUCTURE_REVIEW" in table_block["issues"]
    assert result["pages"][0]["preview_available"] is False


def test_candidate_preview_is_isolated_and_never_sendable(kb):
    client, db, actor = kb
    _, revision = create_text(client, db, text="Обед в столовой с 12:00.")
    page_review(client, revision["id"])
    build = action(client, revision["id"], "candidate-build").json()
    candidate_id = build["candidate_build_result"]["candidate_snapshot_id"]
    employee = save_entity(db, "employee", {"site_id": "a", "employment_stage": "ACTIVE", "employment_type": "OFFICE"})
    conv = save_entity(db, "conversation", {"employee_id": employee.id})
    assert retrieval.answer_question(db, conv, "Обед")["mode"] == "HANDOFF"
    preview = retrieval.preview_answer(db, actor, employee, "Обед", candidate_id)
    assert preview["mode"] == "EXTRACT"
    assert preview["delivery_context"] == "EDITOR_PREVIEW"
    assert preview["provenance"]["publication"] == "CANDIDATE"
    assert not retrieval.answer_still_authorized(db, conv, preview)
    assert retrieval.answer_question(db, conv, "Обед")["mode"] == "HANDOFF"


def test_snapshot_cas_does_not_partially_publish_second_candidate(kb):
    client, db, actor = kb
    first_doc, first_revision = create_text(client, db, "Первое правило про обед.")
    second_doc, second_revision = create_text(client, db, "Второе правило про парковку.")
    for revision in [first_revision, second_revision]:
        page_review(client, revision["id"])
        assert action(client, revision["id"], "candidate-build").status_code == 202
        assert action(client, revision["id"], "qa-runs").status_code == 202
        assert action(client, revision["id"], "review", {"action": "APPROVE", "reason": "Проверено"}).status_code == 200
    assert action(client, first_revision["id"], "publish").status_code == 202
    denied = action(client, second_revision["id"], "publish")
    assert denied.status_code == 409
    assert client.get(f"/knowledge/documents/{second_doc['id']}").json()["active_revision_id"] is None
    assert client.get(f"/knowledge/documents/{first_doc['id']}").json()["active_revision_id"] == first_revision["id"]


def test_scanned_pdf_uses_real_render_and_ocr_route(tmp_path, monkeypatch):
    from PIL import Image, ImageDraw
    from korsik import parsing
    image = Image.new("RGB", (500, 700), "white")
    ImageDraw.Draw(image).text((30, 30), "SCAN: SAMPLE SOURCE", fill="black")
    source = tmp_path / "scan.pdf"
    image.save(source, "PDF", resolution=150)
    calls = []
    def ocr(render, language="rus+eng"):
        calls.append((render.width, render.height, language))
        return "Правило обеденного перерыва", 88.5
    monkeypatch.setattr(parsing, "_ocr", ocr)
    result = parsing.parse_document(source, tmp_path / "previews")
    assert len(calls) == 1
    assert result["blocks"][0]["origin"] == "TESSERACT"
    assert result["pages"][0]["render_sha256"]
    assert (tmp_path / "previews" / result["pages"][0]["preview_file"]).is_file()
    assert result["pages"][0]["review_status"] == "PENDING"


def test_site_hr_cannot_create_or_edit_group_scope(kb):
    client, db, actor = kb
    global_document = client.post("/knowledge/documents", json={"title": "Групповая политика", "scope": {}}).json()
    actor.role, actor.scope = "HR_SPECIALIST", {"site_ids": ["a"]}
    assert client.post("/knowledge/documents", json={"title": "Группа", "scope": {}}).status_code == 403
    assert client.patch(f"/knowledge/documents/{global_document['id']}", json={"title": "Изменено"}, headers={"If-Match": global_document["etag"]}).status_code == 404
    allowed = client.post("/knowledge/documents", json={"title": "Моя площадка", "scope": {"site_ids": ["a"]}})
    assert allowed.status_code == 201


def test_candidate_build_is_durable_and_cancelled_index_never_materializes(kb, monkeypatch):
    client, db, _ = kb
    _, revision = create_text(client, db)
    current = client.get(f"/knowledge/revisions/{revision['id']}")
    original = knowledge.index_chunks
    called = []
    def slow_index(db, chunks, *, checkpoint=None):
        called.append(len(chunks))
        operation = db.scalar(select(Entity).where(Entity.kind == "operation", Entity.data["type"].as_string() == "KNOWLEDGE_CANDIDATE_BUILD").order_by(Entity.created_at.desc()))
        operation.status = "CANCELLED"
        knowledge._update(operation, cancel_requested=True)
        db.commit()
        checkpoint()
        raise AssertionError("cancelled index must stop")
    monkeypatch.setattr(knowledge, "index_chunks", slow_index)
    queued = client.post(f"/knowledge/revisions/{revision['id']}/candidate-build", json={}, headers={"If-Match": current.headers['etag']})
    assert queued.status_code == 202 and queued.json()['status'] == 'PENDING'
    assert called == []
    row = db.get(Entity, revision['id'])
    assert row.data['latest_candidate_operation_id'] == queued.json()['id']
    job = db.scalar(select(Job).where(Job.kind == 'KNOWLEDGE_CANDIDATE_BUILD'))
    knowledge.process_ingest_job(db, job)
    assert called and db.get(Entity, queued.json()['id']).status == 'CANCELLED'
    assert db.scalar(select(Entity).where(Entity.kind == 'knowledge_candidate')) is None
    assert db.scalar(select(Entity).where(Entity.kind == 'knowledge_chunk')) is None
    assert db.get(Entity, revision['id']).data.get('candidate_snapshot_id') is None
    monkeypatch.setattr(knowledge, "index_chunks", original)
    retry = action(client, revision['id'], 'candidate-build').json()
    assert retry['status'] == 'SUCCEEDED'
    assert db.get(Entity, revision['id']).data['candidate_snapshot_id'] == retry['candidate_build_result']['candidate_snapshot_id']


def test_parser_retry_does_not_parse_retained_physical_slides(tmp_path, monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches
    import shutil
    monkeypatch.setattr(shutil, 'which', lambda name: None)
    presentation = Presentation()
    for text in ['Сохранённый текст первой страницы', 'Новая вторая страница']:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = text
    source = tmp_path/'retry.pptx'; presentation.save(source)
    result = parse_document(source, tmp_path/'previews', skip_pages=[1])
    assert len(result['pages']) == 2
    assert result['pages'][0]['status'] == 'RETAINED'
    assert all(b['page_no'] == 2 for b in result['blocks'])
    assert any('Новая вторая' in b.get('text', '') for b in result['blocks'])
    assert not list((tmp_path/'previews').glob('page-0001.png'))


def test_approved_instruction_correction_updates_canonical_card_and_invalidates_qa(kb):
    client, db, _ = kb
    doc = client.post('/knowledge/documents', json={'title':'Проверенная инструкция', 'scope':{'site_ids':['a']}, 'document_category':'INSTRUCTION'}).json()
    created = client.post(f"/knowledge/documents/{doc['id']}/revisions/article", headers={'If-Match':doc['etag']}, json={
        'article_kind':'APPROVED_INSTRUCTION', 'instruction':{'question':'Когда обед?', 'answer_text':'Старый ответ'},
        'dependencies':[{'source_revision_id':'not-published-yet'}]}).json()
    row = db.get(Entity, created['id'])
    knowledge._update(row, qa_artifact_id='old-qa', candidate_snapshot_id='old-candidate', review_fingerprint='old-review'); db.commit()
    extraction = client.get(f"/knowledge/revisions/{row.id}/extraction")
    block_id = extraction.json()['blocks'][0]['id']
    response = client.patch(f'/knowledge/revisions/{row.id}/extraction/blocks/{block_id}', headers={'If-Match':extraction.headers['etag']}, json={
        'patch_type':'TEXT', 'text':'Обед с 12:00 до 13:00.', 'reason':'Сверено с регламентом'})
    assert response.status_code == 200
    current = client.get(f'/knowledge/revisions/{row.id}').json()
    assert current['instruction']['answer_text'] == 'Обед с 12:00 до 13:00.'
    assert not current.get('qa_artifact_id') and not current.get('candidate_snapshot_id') and not current.get('review_fingerprint')


def test_candidate_lease_replacement_discards_uncommitted_chunks(kb, monkeypatch):
    from datetime import timedelta
    client, db, _ = kb
    _, revision = create_text(client, db)
    current = client.get(f"/knowledge/revisions/{revision['id']}")
    queued = client.post(f"/knowledge/revisions/{revision['id']}/candidate-build", json={}, headers={'If-Match':current.headers['etag']})
    job = db.scalar(select(Job).where(Job.kind == 'KNOWLEDGE_CANDIDATE_BUILD'))
    job.payload = {**job.payload, '_lease_token':'old'}
    job.status = 'PROCESSING'; job.lease_until = knowledge.now()+timedelta(seconds=300); db.commit()
    original = knowledge.index_chunks
    def supersede(db, chunks, *, checkpoint=None):
        job.payload = {**job.payload, '_lease_token':'new'}; db.commit()
        return {'mode':'LEXICAL_ONLY', 'indexed':len(chunks)}
    monkeypatch.setattr(knowledge, 'index_chunks', supersede)
    knowledge.process_ingest_job(db, job)
    assert db.get(Entity, queued.json()['id']).status == 'RUNNING'
    assert db.scalar(select(Entity).where(Entity.kind == 'knowledge_candidate')) is None
    assert db.scalar(select(Entity).where(Entity.kind == 'knowledge_chunk')) is None
    monkeypatch.setattr(knowledge, 'index_chunks', original)
    knowledge.process_ingest_job(db, job)
    assert db.get(Entity, queued.json()['id']).status == 'SUCCEEDED'


def test_selected_page_retry_keeps_other_pages_and_human_correction(kb, monkeypatch):
    client, db, _ = kb
    def initial(path, output, checkpoint):
        checkpoint([{'page_no':n,'status':'PROCESSED','review_status':'PENDING','issues':[]} for n in [1,2]],
            [block('TEXT','Старый OCR',1),block('TEXT','Вторая неизменная страница',2)])
        return {'parser_version':'fixture'}
    monkeypatch.setattr(knowledge,'_run_parser',initial)
    doc, revision = create_text(client,db)
    old = client.get(f"/knowledge/revisions/{revision['id']}/extraction")
    first, second = old.json()['blocks']
    corrected = client.patch(f"/knowledge/revisions/{revision['id']}/extraction/blocks/{first['id']}",headers={'If-Match':old.headers['etag']},json={'patch_type':'TEXT','text':'Проверенный текст HR','reason':'Исправлена ошибка OCR'})
    assert corrected.status_code == 200
    def retried(path, output, checkpoint, *, skip_pages=None):
        assert skip_pages == [2]
        checkpoint([{'page_no':1,'status':'PROCESSED','review_status':'PENDING','issues':[]},{'page_no':2,'status':'RETAINED','review_status':'PENDING','issues':[]}], [block('TEXT','Новый машинный OCR',1)])
        return {'parser_version':'fixture'}
    monkeypatch.setattr(knowledge,'_run_parser',retried)
    queued = action(client,revision['id'],'retry',{'page_numbers':[1],'reason':'Повторить распознавание первой страницы'})
    assert queued.status_code == 202
    job = db.scalar(select(Job).where(Job.kind=='KNOWLEDGE_INGEST').order_by(Job.created_at.desc()))
    knowledge.process_ingest_job(db,job)
    latest = client.get(f"/knowledge/revisions/{revision['id']}/extraction").json()
    assert {b['id']:b['text'] for b in latest['blocks']} == {first['id']:'Проверенный текст HR',second['id']:'Вторая неизменная страница'}
    assert 'MANUAL_CORRECTIONS_PRESERVED_REVIEW' in latest['pages'][0]['issues']
    assert latest['pages'][1]['status']=='PROCESSED'
    listing=client.get('/knowledge/documents').json()['items'][0]
    assert listing['latest_revision_status']=='REVIEW_REQUIRED' and listing['latest_ingest_operation_id']==queued.json()['id']
    assert action(client,revision['id'],'retry',{'page_numbers':[1],'page_options':{'1':{'rotation':45}},'reason':'Некорректный угол'}).status_code==422


def test_image_rotation_is_applied_before_ocr_and_preview(tmp_path,monkeypatch):
    from PIL import Image
    from korsik import parsing
    source=tmp_path/'sideways.png';Image.new('RGB',(30,90),'white').save(source)
    seen=[]
    monkeypatch.setattr(parsing,'_ocr',lambda image: (seen.append(image.size) or ('Распознано',99)))
    result=parsing.parse_document(source,tmp_path/'preview',page_options={'1':{'rotation':90,'force_ocr':True}})
    assert seen==[(90,30)]
    assert result['pages'][0]['rotation_degrees']==90
    assert result['blocks'][0]['text']=='Распознано'
    with Image.open(tmp_path/'preview'/'page-0001.png') as preview:assert preview.width>preview.height


def test_published_impact_review_stays_readable_and_immutable(kb):
    client,db,_=kb
    _,revision=create_text(client,db);publish(client,revision['id'])
    impact_id=db.get(Entity,revision['id']).data['impact_review_id']
    old=client.get(f'/knowledge/impact-reviews/{impact_id}')
    assert old.status_code==200 and old.headers['etag']
    result=client.post(f'/knowledge/impact-reviews/{impact_id}/review',json={'decisions':[]},headers={'If-Match':old.headers['etag']})
    assert result.status_code==409
    assert client.get(f'/knowledge/impact-reviews/{impact_id}').headers['etag']==old.headers['etag']
