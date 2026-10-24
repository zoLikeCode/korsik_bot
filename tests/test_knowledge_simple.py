"""Opt-in preparation is automatic/private; publication is exact explicit approval."""
from copy import deepcopy
from datetime import timedelta
import pytest
from sqlalchemy import select
from korsik import knowledge as kb, knowledge_simple as simple, retrieval, knowledge_admin
from korsik.common import save_entity
from korsik.db import now
from korsik.models import Entity, Job
from korsik.parsing import block
from test_knowledge import kb as kb_fixture


@pytest.fixture
def flow(kb_fixture):
    client, db, actor = kb_fixture
    client.app.include_router(simple.router)
    return client, db, actor


def upload(flow, text="Правила передачи документов: пароль отправляется отдельным сообщением.", doc=None):
    client, db, _ = flow
    doc = doc or client.post('/knowledge/documents', json={'title':'Правила','scope':{'site_ids':['a'],'employment_stages':['ACTIVE']},'document_category':'REGULATION'}).json()
    fresh=client.get('/knowledge/documents/'+doc['id'])
    result=client.post('/knowledge/documents/'+doc['id']+'/revisions',files={'file':('rules.txt',text.encode(),'text/plain')},data={'simple_workflow':'true'},headers={'If-Match':fresh.headers['etag']})
    assert result.status_code==202,result.text
    return doc,result.json()['revision']


def drain(db):
    for _ in range(25):
        job=db.scalar(select(Job).where(Job.status.in_(['PENDING','RETRY'])).order_by(Job.created_at,Job.id))
        if not job: return
        token='synthetic-simple-'+job.id
        job.status='PROCESSING';job.lease_until=now()+timedelta(minutes=10);job.payload={**job.payload,'_lease_token':token};db.commit()
        assert kb.process_ingest_job(db,job)
        db.refresh(job)
        if job.payload.get('_lease_token')==token: job.status='COMPLETED'
        db.commit()
    raise AssertionError('automatic preparation loop')


def status(flow,rid): return flow[0].get(f'/knowledge/revisions/{rid}/simple-status')
def publish(flow,rid,etag=None):
    client=flow[0];current=client.get(f'/knowledge/revisions/{rid}')
    return client.post(f'/knowledge/revisions/{rid}/simple-publish',json={},headers={'If-Match':etag or current.headers['etag']})


def test_native_text_automatic_private_pipeline_and_single_explicit_publish(flow):
    client,db,actor=flow
    doc,rev=upload(flow);drain(db)
    row=db.get(Entity,rev['id']);ext=kb._extraction(db,row)
    assert all(p.get('review_status')!='REVIEWED' for p in ext.data['pages'])
    assert row.status=='REVIEW_REQUIRED'
    assert status(flow,row.id).json()['ready_to_publish'] is True
    person=save_entity(db,'employee',{'site_id':'a','employment_stage':'ACTIVE','employment_type':'OFFICE'})
    assert retrieval.published_corpus(db,person)=={}
    assert not db.get(Entity,doc['id']).data.get('active_revision_id')
    result=publish(flow,row.id)
    assert result.status_code==200,result.text
    assert result.json()['status']=='PUBLISHED'
    assert db.get(Entity,doc['id']).data['active_revision_id']==row.id
    assert row.id in retrieval.published_corpus(db,person)
    assert row.data['simple_approval']['method']=='EXPLICIT_WHOLE_REVISION_APPROVAL'
    assert all(p.get('review_status')!='REVIEWED' for p in ext.data['pages'])


def test_status_GET_does_not_enqueue_mutate_or_publish(flow):
    client,db,_=flow
    _,rev=upload(flow);drain(db)
    before=[(e.id,e.version,e.status,deepcopy(e.data)) for e in db.scalars(select(Entity))]
    jobs=len(list(db.scalars(select(Job))))
    assert status(flow,rev['id']).status_code==200
    assert status(flow,rev['id']).status_code==200
    assert before==[(e.id,e.version,e.status,e.data) for e in db.scalars(select(Entity))]
    assert len(list(db.scalars(select(Job))))==jobs


def changed_extraction(flow,rid,payload):
    db=flow[1];row=db.get(Entity,rid);kb._new_extraction(db,row,payload);db.commit();drain(db)
    return row,kb._extraction(db,row)


def test_undecoded_figures_are_reference_only_but_pure_image_requires_description(flow):
    _,db,_=flow
    _,rev=upload(flow);drain(db)
    image=block('FIGURE','',1,image_sha256='a'*64)
    row,ext=changed_extraction(flow,rev['id'],{'pages':[{'page_no':1,'status':'PROCESSED','review_status':'PENDING','issues':['FIGURE_DESCRIPTION_REQUIRED']}], 'blocks':[block('TEXT','Достоверное содержимое регламента',1),image]})
    figure=next(b for b in ext.data['blocks'] if b['kind']=='FIGURE')
    assert figure['disposition']=='REFERENCE_ONLY'
    assert figure.get('description') in (None,'')
    assert status(flow,row.id).json()['ready_to_publish'] is True
    row,_=changed_extraction(flow,row.id,{'pages':[{'page_no':1,'status':'PROCESSED','review_status':'PENDING','issues':['FIGURE_DESCRIPTION_REQUIRED']}],'blocks':[image]})
    assert status(flow,row.id).json()['ready_to_publish'] is False
    assert publish(flow,row.id).status_code==409


@pytest.mark.parametrize('issue',['OCR_LOW_CONFIDENCE','TABLE_STRUCTURE_UNCERTAIN','READING_ORDER_UNCERTAIN'])
def test_dangerous_parser_issues_block_even_explicit_publish(flow,issue):
    _,rev=upload(flow);drain(flow[1])
    changed_extraction(flow,rev['id'],{'pages':[{'page_no':1,'status':'PROCESSED','review_status':'PENDING','issues':[issue]}],'blocks':[block('TEXT','Распознанный текст требует сверки',1)]})
    result=status(flow,rev['id']).json()
    assert result['ready_to_publish'] is False
    assert issue in str(result['blockers'])
    assert publish(flow,rev['id']).status_code==409


def test_control_FAIL_cannot_be_approved_by_single_click(flow,monkeypatch):
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',lambda *a,**k:{'status':'FAILED','results':[],'fingerprint':'failed-machine-suite'})
    doc,rev=upload(flow);drain(flow[1])
    assert status(flow,rev['id']).json()['ready_to_publish'] is False
    assert publish(flow,rev['id']).status_code==409
    assert not flow[1].get(Entity,doc['id']).data.get('active_revision_id')


def test_machine_passed_controls_whole_version_approval_is_truthful(flow,monkeypatch):
    monkeypatch.setattr(knowledge_admin,'evaluate_control_suite',lambda *a,**k:{'status':'PASSED','results':[{'status':'PASSED','question_version_id':'synthetic'}],'fingerprint':'passed-machine-suite'})
    _,rev=upload(flow);drain(flow[1])
    row=flow[1].get(Entity,rev['id']);qa=flow[1].get(Entity,row.data['qa_artifact_id'])
    assert qa.data['manual_review_status']=='PENDING'
    assert publish(flow,row.id).status_code==200
    assert qa.data['manual_review_status']=='PENDING'  # no forged per-answer PASS
    assert row.data['simple_approval']['qa_id']==qa.id


def test_changed_content_and_stale_etag_cannot_publish_unseen_revision(flow):
    client,db,_=flow
    _,rev=upload(flow);drain(db)
    shown=client.get(f"/knowledge/revisions/{rev['id']}")
    row=db.get(Entity,rev['id']);payload=deepcopy(kb._extraction(db,row).data);payload['blocks'][0]['text']='Новая редакция — другой текст'
    kb._new_extraction(db,row,payload);db.commit();drain(db)
    assert publish(flow,row.id,shown.headers['etag']).status_code==412
    assert row.status!='PUBLISHED'


def test_old_published_version_remains_active_until_successful_replacement(flow):
    doc,first=upload(flow);drain(flow[1]);assert publish(flow,first['id']).status_code==200
    _,second=upload(flow,'Заменяющий документ с новым порядком передачи файлов.',doc=doc)
    assert flow[1].get(Entity,doc['id']).data['active_revision_id']==first['id']
    drain(flow[1]);assert flow[1].get(Entity,doc['id']).data['active_revision_id']==first['id']
    assert publish(flow,second['id']).status_code==200
    assert flow[1].get(Entity,doc['id']).data['active_revision_id']==second['id']


def test_stale_head_and_revoked_HR_authority_block_publication(flow):
    _,db,actor=flow
    _,rev=upload(flow);drain(db)
    head=kb._head(db);old=head.data['snapshot_id'];snapshot=save_entity(db,'knowledge_snapshot',{'members':{}},status='PUBLISHED')
    kb._update(head,snapshot_id=snapshot.id);db.commit()
    assert publish(flow,rev['id']).status_code==409
    assert status(flow,rev['id']).json()['requires_prepare'] is True
    kb._update(head,snapshot_id=old);db.commit()
    actor.role='HR_SPECIALIST'
    assert publish(flow,rev['id']).status_code==403


def test_pure_image_HR_description_becomes_searchable_and_publishable(flow):
    client,db,_=flow
    _,rev=upload(flow);drain(db)
    image=block('FIGURE','',1,image_sha256='b'*64)
    row,ext=changed_extraction(flow,rev['id'],{'pages':[{'page_no':1,'status':'PROCESSED','review_status':'PENDING','issues':['FIGURE_DESCRIPTION_REQUIRED']}],'blocks':[image]})
    assert status(flow,row.id).json()['ready_to_publish'] is False
    current=client.get(f'/knowledge/revisions/{row.id}/extraction')
    result=client.patch(f"/knowledge/revisions/{row.id}/extraction/blocks/{image['id']}",json={'patch_type':'FIGURE_DESCRIPTION','description':'План прохода в офис: вход со стороны центральной проходной.','reason':'Проверено по изображению','expected_extraction_revision_id':ext.id},headers={'If-Match':current.headers['etag']})
    assert result.status_code==200,result.text
    drain(db)
    latest=kb._extraction(db,row)
    assert latest.data['blocks'][0]['disposition']=='INCLUDE'
    assert status(flow,row.id).json()['ready_to_publish'] is True
    assert publish(flow,row.id).status_code==200


def test_failed_publication_rolls_back_whole_approval_and_old_active(flow,monkeypatch):
    doc,first=upload(flow);drain(flow[1]);assert publish(flow,first['id']).status_code==200
    _,second=upload(flow,'Обновлённый порядок работы с документами.',doc=doc);drain(flow[1])
    original=kb._publish
    def fault(*args,**kwargs):
        original(*args,**kwargs)
        from fastapi import HTTPException
        raise HTTPException(409,'synthetic_after_atomic_write')
    monkeypatch.setattr(kb,'_publish',fault)
    assert publish(flow,second['id']).status_code==409
    flow[1].expire_all()
    assert flow[1].get(Entity,doc['id']).data['active_revision_id']==first['id']
    row=flow[1].get(Entity,second['id'])
    assert row.status=='REVIEW_REQUIRED'
    assert row.data.get('simple_approval') is None


def test_existing_advanced_draft_explicit_prepare_optin_without_mutating_GET(flow):
    from test_knowledge import create_text
    client,db,_=flow
    _,revision=create_text(client,db)
    before=len(list(db.scalars(select(Job))))
    assert status(flow,revision['id']).json()['requires_prepare'] is True
    assert len(list(db.scalars(select(Job))))==before
    current=client.get(f"/knowledge/revisions/{revision['id']}")
    result=client.post(f"/knowledge/revisions/{revision['id']}/simple-prepare",json={},headers={'If-Match':current.headers['etag']})
    assert result.status_code==202,result.text
    assert result.headers.get('etag')
    drain(db)
    assert status(flow,revision['id']).json()['ready_to_publish'] is True
    assert publish(flow,revision['id']).status_code==200


def test_real_native_PPTX_auto_prepares_without_page_review_clicks(flow,tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    presentation=Presentation()
    for index in range(2):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        shape=slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2))
        shape.text_frame.text=f'Регламент. Раздел {index+1}. Документы передаются через защищённый канал компании.'
    image=tmp_path/'source-figure.png';Image.new('RGB',(100,100),'navy').save(image)
    slide.shapes.add_picture(str(image),Inches(1),Inches(4),width=Inches(1))
    path=tmp_path/'policy.pptx';presentation.save(path)
    client,db,_=flow
    doc=client.post('/knowledge/documents',json={'title':'Презентация','scope':{'site_ids':['a'],'employment_stages':['ACTIVE']},'document_category':'OVERVIEW'}).json()
    result=client.post(f"/knowledge/documents/{doc['id']}/revisions",files={'file':('policy.pptx',path.read_bytes(),'application/vnd.openxmlformats-officedocument.presentationml.presentation')},data={'simple_workflow':'true'},headers={'If-Match':doc['etag']})
    assert result.status_code==202,result.text
    rid=result.json()['revision']['id'];drain(db)
    row=db.get(Entity,rid);ext=kb._extraction(db,row)
    assert len(ext.data['pages'])==2
    assert any(b.get('kind')=='FIGURE' and b.get('disposition')=='REFERENCE_ONLY' for b in ext.data['blocks'])
    assert all(p.get('review_status')!='REVIEWED' for p in ext.data['pages'])
    assert status(flow,rid).json()['ready_to_publish'] is True,status(flow,rid).json()
    assert publish(flow,rid).status_code==200


def test_stale_principal_scope_source_and_disabled_account_fail_closed(flow):
    from korsik.models import User
    client,db,actor=flow
    _,rev=upload(flow);drain(db)
    user=db.get(User,actor.id);user.enabled=False;db.commit()
    assert publish(flow,rev['id']).status_code==403
    user.enabled=True;user.scope={'site_ids':['other']};db.commit()
    assert publish(flow,rev['id']).status_code==404
    user.scope={'all':True};db.commit()
    row=db.get(Entity,rev['id'])
    from pathlib import Path
    Path(row.data['source_path']).write_text('Изменение источника после проверки')
    assert publish(flow,row.id).status_code==409
    assert row.status=='REVIEW_REQUIRED'


def test_old_APPROVED_can_explicitly_opt_in_and_rebuild_without_forged_review(flow):
    from test_knowledge import create_text,page_review,action
    client,db,_=flow
    _,rev=create_text(client,db);page_review(client,rev['id']);action(client,rev['id'],'candidate-build');action(client,rev['id'],'qa-runs')
    assert action(client,rev['id'],'review',{'action':'APPROVE','reason':'Ранее проверено HR'}).status_code==200
    assert status(flow,rev['id']).json()['requires_prepare'] is True
    current=client.get(f"/knowledge/revisions/{rev['id']}")
    prepared=client.post(f"/knowledge/revisions/{rev['id']}/simple-prepare",json={},headers={'If-Match':current.headers['etag']})
    assert prepared.status_code==202,prepared.text
    row=db.get(Entity,rev['id']);assert row.status=='REVIEW_REQUIRED' and row.data.get('review_fingerprint') is None
    drain(db);assert publish(flow,row.id).status_code==200


def test_failed_index_is_recoverable_without_busy_loop_or_silent_publication(flow,monkeypatch):
    original=kb.index_chunks
    monkeypatch.setattr(kb,'index_chunks',lambda *a,**k:(_ for _ in ()).throw(RuntimeError('synthetic_index_failure')))
    client,db,_=flow
    doc,rev=upload(flow);drain(db)
    result=status(flow,rev['id']).json()
    assert result['processing'] is False and result['requires_prepare'] is True
    assert result['ready_to_publish'] is False
    assert not db.get(Entity,doc['id']).data.get('active_revision_id')
    monkeypatch.setattr(kb,'index_chunks',original)
    current=client.get(f"/knowledge/revisions/{rev['id']}")
    assert client.post(f"/knowledge/revisions/{rev['id']}/simple-prepare",json={},headers={'If-Match':current.headers['etag']}).status_code==202
    drain(db);assert status(flow,rev['id']).json()['ready_to_publish'] is True
    assert publish(flow,rev['id']).status_code==200


def test_changed_extraction_during_index_discards_old_candidate_and_reprepares(flow,monkeypatch):
    original=kb.index_chunks;changed=[False]
    client,db,_=flow
    doc,rev=upload(flow)
    def change_once(db,chunks,checkpoint=None):
        result=original(db,chunks,checkpoint=checkpoint)
        if not changed[0]:
            changed[0]=True
            row=db.get(Entity,rev['id']);payload=deepcopy(kb._extraction(db,row).data)
            payload['blocks'][0]['text']='Новый проверяемый порядок передачи: пароли отправляются отдельно.'
            kb._new_extraction(db,row,payload);db.commit()
        return result
    monkeypatch.setattr(kb,'index_chunks',change_once)
    drain(db)
    row=db.get(Entity,rev['id']);candidate=db.get(Entity,row.data['candidate_snapshot_id'])
    assert changed[0] and status(flow,row.id).json()['ready_to_publish'] is True
    assert candidate.data['members'][row.id]['extraction_revision_id']==row.data['latest_extraction_revision_id']
    assert row.status=='REVIEW_REQUIRED' and not db.get(Entity,doc['id']).data.get('active_revision_id')
    assert all(chunk.status!='PUBLISHED' for chunk in db.scalars(select(Entity).where(Entity.kind=='knowledge_chunk')))


def test_expired_coordinator_lease_cannot_queue_candidate(flow):
    _,db,_=flow
    _,rev=upload(flow)
    ingest=db.scalar(select(Job).where(Job.kind=='KNOWLEDGE_INGEST'))
    assert kb.process_ingest_job(db,ingest)
    ingest.status='COMPLETED';db.commit()
    job=db.scalar(select(Job).where(Job.kind==simple.KIND))
    job.status='PROCESSING';job.payload={**job.payload,'_lease_token':'expired'};job.lease_until=now()-timedelta(seconds=1);db.commit()
    with pytest.raises(kb.IngestLeaseLost): kb.process_ingest_job(db,job)
    db.rollback()
    assert db.get(Entity,rev['id']).data.get('candidate_snapshot_id') is None
    assert not list(db.scalars(select(Job).where(Job.kind=='KNOWLEDGE_CANDIDATE_BUILD')))
